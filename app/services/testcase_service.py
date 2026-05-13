import json
import uuid
import re
from typing import List, Dict, Any, Optional
from datetime import datetime
from app.models.schemas import TestCaseItem
from app.services.llm_service import llm_service
from app.services.document_service import document_reader
from app.services.document_manager import uploaded_files
from app.core.config import settings
from app.core.logger import get_logger
import os

logger = get_logger("testcase_service")


class TestCaseService:
    def __init__(self):
        self.testcases: Dict[str, List[TestCaseItem]] = {}

    async def generate_testcases(
        self,
        document_ids: List[str],
        endpoint_url: str,
        method: str,
        count: int = 10,
    ) -> List[TestCaseItem]:
        logger.info(f"Generating {count} testcases for {method} {endpoint_url}")

        context = self._prepare_context(document_ids)

        prompt = self._build_generation_prompt(
            endpoint_url=endpoint_url,
            method=method,
            count=count,
            context=context,
        )

        try:
            response = await llm_service.llm.ainvoke(prompt)
            response_text = (
                response.content if hasattr(response, "content") else str(response)
            )

            testcases = self._parse_llm_response(response_text, endpoint_url, method)
            testcase_id = str(uuid.uuid4())
            self.testcases[testcase_id] = testcases

            logger.info(f"Generated {len(testcases)} testcases, id: {testcase_id}")
            return testcases
        except Exception as e:
            logger.error(f"Failed to generate testcases: {e}")
            raise

    def _prepare_context(self, document_ids: List[str]) -> str:
        if not document_ids:
            return ""

        file_paths = []
        for doc in uploaded_files:
            if doc.filename in document_ids:
                file_path = os.path.join(settings.UPLOAD_DIR, doc.filename)
                if os.path.exists(file_path):
                    file_paths.append(file_path)

        if file_paths:
            return document_reader.read_all_documents(file_paths)
        return ""

    def _build_generation_prompt(
        self, endpoint_url: str, method: str, count: int, context: str
    ) -> str:
        context_section = (
            f"""
接口文档内容:
{context}
"""
            if context
            else ""
        )

        prompt = f"""你是一个专业的接口测试工程师。请根据以下信息生成{count}个测试用例。

接口信息:
- 接口地址: {endpoint_url}
- 请求方法: {method}
{context_section}

请生成测试用例，包含以下字段:
1. id: 用例编号 (从1开始)
2. name: 用例名称 (简洁描述测试场景)
3. description: 用例描述
4. method: 请求方法
5. url: 请求URL
6. headers: 请求头 (如需要)
7. params: URL参数 (如需要)
8. body: 请求体 (如需要，JSON格式)
9. expected_status: 预期响应状态码
10. priority: 优先级 (P0/P1/P2/P3)

要求:
- 覆盖正常场景、边界条件、异常场景
- P0: 核心功能测试，必须通过
- P1: 重要功能测试
- P2: 一般功能测试
- P3: 边缘场景测试
- 返回JSON数组格式，不要有其他文字"""

        return prompt

    def _parse_llm_response(
        self, response_text: str, endpoint_url: str, method: str
    ) -> List[TestCaseItem]:
        json_match = re.search(r"\[.*\]", response_text, re.DOTALL)
        if not json_match:
            logger.info(
                "No JSON array found in LLM response, creating default testcases"
            )
            return self._create_default_testcases(endpoint_url, method)

        try:
            data = json.loads(json_match.group())
            testcases = []
            for i, item in enumerate(data):
                testcase = TestCaseItem(
                    id=item.get("id", i + 1),
                    name=item.get("name", f"测试用例 {i + 1}"),
                    description=item.get("description"),
                    method=item.get("method", method),
                    url=item.get("url", endpoint_url),
                    headers=item.get("headers", {}),
                    params=item.get("params"),
                    body=item.get("body"),
                    expected_status=item.get("expected_status", 200),
                    priority=item.get("priority", "P1"),
                )
                testcases.append(testcase)
            return testcases
        except json.JSONDecodeError as e:
            logger.info(f"Failed to parse JSON: {e}, using defaults")
            return self._create_default_testcases(endpoint_url, method)

    def _create_default_testcases(
        self, endpoint_url: str, method: str
    ) -> List[TestCaseItem]:
        defaults = [
            TestCaseItem(
                id=1,
                name="正常请求",
                description="验证接口正常调用",
                method=method,
                url=endpoint_url,
                expected_status=200,
                priority="P0",
            ),
            TestCaseItem(
                id=2,
                name="空参数",
                description="验证空参数处理",
                method=method,
                url=endpoint_url,
                params={},
                expected_status=400,
                priority="P1",
            ),
            TestCaseItem(
                id=3,
                name="缺少必填参数",
                description="验证必填参数缺失时的行为",
                method=method,
                url=endpoint_url,
                expected_status=400,
                priority="P1",
            ),
        ]
        return defaults

    def export_json(self, testcases: List[TestCaseItem]) -> str:
        data = [
            {
                "id": tc.id,
                "name": tc.name,
                "description": tc.description,
                "method": tc.method,
                "url": tc.url,
                "headers": tc.headers,
                "params": tc.params,
                "body": tc.body,
                "expected_status": tc.expected_status,
                "priority": tc.priority,
            }
            for tc in testcases
        ]
        return json.dumps(data, ensure_ascii=False, indent=2)

    def export_postman(
        self,
        testcases: List[TestCaseItem],
        collection_name: str = "API Test Collection",
    ) -> dict:
        items = []
        for tc in testcases:
            item = {
                "name": tc.name,
                "request": {
                    "method": tc.method,
                    "header": [
                        {"key": k, "value": v} for k, v in (tc.headers or {}).items()
                    ],
                    "url": {
                        "raw": tc.url,
                        "protocol": "https",
                        "host": self._extract_host(tc.url),
                        "path": self._extract_path(tc.url),
                        "query": [
                            {"key": k, "value": str(v)}
                            for k, v in (tc.params or {}).items()
                        ],
                    },
                },
            }
            if tc.body and tc.method in ["POST", "PUT", "PATCH"]:
                item["request"]["body"] = {
                    "mode": "raw",
                    "raw": json.dumps(tc.body),
                    "options": {"raw": {"language": "json"}},
                }
            items.append(item)

        collection = {
            "info": {
                "name": collection_name,
                "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
            },
            "item": items,
        }
        return collection

    def _extract_host(self, url: str) -> List[str]:
        match = re.search(r"https?://([^/]+)", url)
        if match:
            return match.group(1).split(".")
        return ["example", "com"]

    def _extract_path(self, url: str) -> List[str]:
        match = re.search(r"https?://[^/]+(/.*)", url)
        if match:
            path = match.group(1)
            return [p for p in path.split("/") if p]
        return ["api"]

    def export_excel(self, testcases: List[TestCaseItem]) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "测试用例"

            header_fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
            header_font = Font(bold=True, color="FFFFFF")

            headers = [
                "ID",
                "用例名称",
                "描述",
                "方法",
                "URL",
                "请求头",
                "参数",
                "请求体",
                "预期状态",
                "优先级",
            ]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for row, tc in enumerate(testcases, 2):
                ws.cell(row=row, column=1, value=tc.id)
                ws.cell(row=row, column=2, value=tc.name)
                ws.cell(row=row, column=3, value=tc.description or "")
                ws.cell(row=row, column=4, value=tc.method)
                ws.cell(row=row, column=5, value=tc.url)
                ws.cell(
                    row=row,
                    column=6,
                    value=json.dumps(tc.headers, ensure_ascii=False)
                    if tc.headers
                    else "",
                )
                ws.cell(
                    row=row,
                    column=7,
                    value=json.dumps(tc.params, ensure_ascii=False)
                    if tc.params
                    else "",
                )
                ws.cell(
                    row=row,
                    column=8,
                    value=json.dumps(tc.body, ensure_ascii=False) if tc.body else "",
                )
                ws.cell(row=row, column=9, value=tc.expected_status)
                ws.cell(row=row, column=10, value=tc.priority)

            for col in range(1, 11):
                ws.column_dimensions[chr(64 + col)].width = 15

            from io import BytesIO

            output = BytesIO()
            wb.save(output)
            return output.getvalue()
        except ImportError:
            raise ImportError("openpyxl is not installed. Run: pip install openpyxl")

    def get_testcases(self, testcase_id: str) -> Optional[List[TestCaseItem]]:
        return self.testcases.get(testcase_id)


testcase_service = TestCaseService()
